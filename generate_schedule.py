"""
客語堂主日服事表生成器
Usage: python3 generate_schedule.py [config.json]

All constraints are read from config.json (default).
Edit config.json to change the quarter, role pools, or any constraint —
no code changes needed.
"""

import json
import os
import sys
import random
import datetime
from collections import defaultdict
from copy import deepcopy

# ── Force UTF-8 ───────────────────────────────────────────────────────────────
if sys.stdout.encoding and sys.stdout.encoding.lower() != "utf-8":
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")

# ── Auto-install openpyxl ─────────────────────────────────────────────────────
try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.styles.colors import Color
    from openpyxl.utils import get_column_letter
except ImportError:
    import subprocess
    subprocess.run([sys.executable, "-m", "pip", "install",
                    "--break-system-packages", "--user", "openpyxl"], check=True)
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.styles.colors import Color
    from openpyxl.utils import get_column_letter

BASE_DIR = os.path.dirname(os.path.abspath(__file__))

# ─────────────────────────────────────────────────────────────────────────────
# Load config
# ─────────────────────────────────────────────────────────────────────────────

def load_config(path: str) -> dict:
    with open(path, encoding="utf-8") as f:
        return json.load(f)


def quarter_sundays(year: int, quarter: int,
                     include_first_sunday: bool = True) -> list[datetime.date]:
    """Return all Sundays in the given quarter.

    If include_first_sunday is False, the first Sunday of each calendar
    month is left out (e.g. because that Sunday is a combined/Communion
    service with no separate duty roster)."""
    starts = {1: (1, 1), 2: (4, 1), 3: (7, 1), 4: (10, 1)}
    ends   = {1: (3, 31), 2: (6, 30), 3: (9, 30), 4: (12, 31)}
    start  = datetime.date(year, *starts[quarter])
    end    = datetime.date(year, *ends[quarter])
    d      = start + datetime.timedelta(days=(6 - start.weekday()) % 7)
    sundays = []
    while d <= end:
        if include_first_sunday or d.day > 7:
            sundays.append(d)
        d += datetime.timedelta(weeks=1)
    return sundays


# ─────────────────────────────────────────────────────────────────────────────
# Build runtime tables from config
# ─────────────────────────────────────────────────────────────────────────────

def _date_range(start: datetime.date, end: datetime.date) -> list[datetime.date]:
    days = []
    d = start
    while d <= end:
        days.append(d)
        d += datetime.timedelta(days=1)
    return days


def _month_range(start_year: int, start_month: int,
                  end_year: int, end_month: int) -> list[datetime.date]:
    """All calendar days from the 1st of (start_year, start_month) through
    the last day of (end_year, end_month), inclusive."""
    start = datetime.date(start_year, start_month, 1)
    if end_month == 12:
        end = datetime.date(end_year, 12, 31)
    else:
        end = datetime.date(end_year, end_month + 1, 1) - datetime.timedelta(days=1)
    return _date_range(start, end)


def parse_unavailable_entry(token: str) -> list[datetime.date]:
    """Parse one entry from an unavailable_dates list. Supported formats:
      - "YYYY-MM-DD"            a single date
      - "YYYY-MM"               every day in that month
      - "YYYY-MM-DD:YYYY-MM-DD" every date in that (inclusive) range
      - "YYYY-MM:YYYY-MM"       every day in that (inclusive) range of months
    """
    token = token.strip()
    if ":" in token:
        start_s, end_s = (s.strip() for s in token.split(":", 1))
        if len(start_s) == 7 and len(end_s) == 7:
            sy, sm = (int(x) for x in start_s.split("-"))
            ey, em = (int(x) for x in end_s.split("-"))
            return _month_range(sy, sm, ey, em)
        return _date_range(datetime.date.fromisoformat(start_s),
                            datetime.date.fromisoformat(end_s))
    if len(token) == 7:
        y, m = (int(x) for x in token.split("-"))
        return _month_range(y, m, y, m)
    return [datetime.date.fromisoformat(token)]


def build_tables(cfg: dict, sundays: list[datetime.date]):
    c = cfg["constraints"]

    # max_per_role: {(person, role): max_count}
    max_per_role: dict[tuple, int] = {}
    for item in c.get("max_per_role", []):
        max_per_role[(item["person"], item["role"])] = item["max"]

    # available_months: {(person, role): frozenset of months}
    # None means available all season
    avail_months: dict[tuple, frozenset | None] = {}
    for item in c.get("available_months", []):
        key = (item["person"], item["role"])
        months = item.get("months")
        avail_months[key] = frozenset(months) if months else None

    # hard_pairs: {mc_name: frozenset of allowed projection names}
    hard_pairs: dict[str, frozenset] = {}
    for item in c.get("hard_pairs", []):
        hard_pairs[item["mc"]] = frozenset(item["projection_must_be"])

    # no_same_day: list of (person, frozenset of roles)
    no_same_day: list[tuple[str, frozenset]] = []
    for item in c.get("no_same_day", []):
        no_same_day.append((item["person"], frozenset(item["roles"])))

    # unavailable_dates: {person: frozenset of datetime.date}
    unavailable: dict[str, frozenset] = {}
    for item in c.get("unavailable_dates", []):
        person = item["person"]
        dates  = frozenset(
            d for token in item["dates"] for d in parse_unavailable_entry(token)
        )
        unavailable[person] = unavailable.get(person, frozenset()) | dates

    # sunday month index lookup
    sunday_months = [s.month for s in sundays]

    return max_per_role, avail_months, hard_pairs, no_same_day, sunday_months, unavailable


# ─────────────────────────────────────────────────────────────────────────────
# Score helpers
# ─────────────────────────────────────────────────────────────────────────────

def compute_scores(schedule: list[dict], all_people: list[str]) -> dict[str, int]:
    scores: dict[str, int] = {p: 0 for p in all_people}
    for day in schedule:
        for role in ("司會", "司琴", "投影"):
            name = day.get(role)
            if name:
                scores[name] = scores.get(name, 0) + 1
    return scores


def score_variance(scores: dict[str, int]) -> float:
    vals = list(scores.values())
    mean = sum(vals) / len(vals)
    return sum((v - mean) ** 2 for v in vals) / len(vals)


# ─────────────────────────────────────────────────────────────────────────────
# Constraint helpers
# ─────────────────────────────────────────────────────────────────────────────

def get_piano_pool(idx: int, mc: str, role_counts: dict[tuple, int],
                   piano_pool: list[str],
                   max_per_role, avail_months, no_same_day,
                   sunday_months, unavailable, sundays) -> list[str]:
    month  = sunday_months[idx]
    date   = sundays[idx]
    pool = []
    for p in piano_pool:
        # leave of absence
        if date in unavailable.get(p, frozenset()):
            continue
        # max cap
        if max_per_role.get((p, "司琴"), 999) <= role_counts.get((p, "司琴"), 0):
            continue
        # availability window
        key = (p, "司琴")
        if key in avail_months and avail_months[key] is not None:
            if month not in avail_months[key]:
                continue
        # no_same_day
        if any(p == person and "司會" in roles and "司琴" in roles and mc == person
               for person, roles in no_same_day):
            continue
        pool.append(p)
    return pool


def get_proj_pool(idx: int, mc: str, proj_pool: list[str],
                  hard_pairs, no_same_day, unavailable, sundays) -> list[str]:
    required = hard_pairs.get(mc)
    date     = sundays[idx]
    pool = []
    for p in proj_pool:
        # leave of absence
        if date in unavailable.get(p, frozenset()):
            continue
        if any(p == person and "司會" in roles and "投影" in roles and mc == person
               for person, roles in no_same_day):
            continue
        if required is not None and p not in required:
            continue
        pool.append(p)
    return pool


def is_valid_full(schedule: list[dict], cfg_tables, sunday_months, sundays) -> bool:
    max_per_role, avail_months, hard_pairs, no_same_day, _, unavailable = cfg_tables
    role_counts: dict[tuple, int] = defaultdict(int)

    for i, day in enumerate(schedule):
        mc    = day["司會"]
        piano = day["司琴"]
        proj  = day["投影"]
        month = sunday_months[i]
        date  = sundays[i]

        # unavailable_dates
        for name in (mc, piano, proj):
            if date in unavailable.get(name, frozenset()):
                return False

        # max_per_role
        for role, name in (("司會", mc), ("司琴", piano), ("投影", proj)):
            role_counts[(name, role)] += 1
            if role_counts[(name, role)] > max_per_role.get((name, role), 999):
                return False

        # available_months
        key = (piano, "司琴")
        if key in avail_months and avail_months[key] is not None:
            if month not in avail_months[key]:
                return False

        # hard_pairs
        if mc in hard_pairs and proj not in hard_pairs[mc]:
            return False

        # no_same_day
        for person, roles in no_same_day:
            served = {r for r in ("司會", "司琴", "投影")
                      if day[r] == person and r in roles}
            if len(served) == len(roles):
                return False

        # no consecutive same person in same role
        if i > 0:
            prev = schedule[i - 1]
            for role_key in ("司會", "司琴", "投影"):
                if day[role_key] == prev[role_key]:
                    return False

    return True


# ─────────────────────────────────────────────────────────────────────────────
# Greedy + local-swap solver
# ─────────────────────────────────────────────────────────────────────────────

def greedy_pass(seed: int, sundays, mc_pool, piano_pool, proj_pool,
                all_people, cfg_tables, sunday_months) -> list[dict]:
    max_per_role, avail_months, hard_pairs, no_same_day, _, unavailable = cfg_tables
    rng = random.Random(seed)
    schedule: list[dict] = []
    scores: dict[str, int] = {p: 0 for p in all_people}
    role_counts: dict[tuple, int] = defaultdict(int)

    def sort_key(name, last_name=None):
        # Penalise serving the same role two Sundays in a row (soft: still
        # falls back to them if no one else is available for that role)
        consec = 1 if name == last_name else 0
        return (consec, scores[name], rng.random())

    for i, sunday in enumerate(sundays):
        last = schedule[-1] if schedule else {}

        # 司會 — also skip if on leave
        mc = None
        for candidate in sorted(mc_pool,
                                 key=lambda n: sort_key(n, last.get("司會"))):
            if sunday in unavailable.get(candidate, frozenset()):
                continue
            if get_proj_pool(i, candidate, proj_pool, hard_pairs, no_same_day,
                             unavailable, sundays):
                mc = candidate
                break
        if mc is None:
            raise RuntimeError(f"No valid 司會 for Sunday {i} ({sunday})")

        # 司琴
        piano_cands = sorted(
            get_piano_pool(i, mc, role_counts, piano_pool,
                           max_per_role, avail_months, no_same_day,
                           sunday_months, unavailable, sundays),
            key=lambda n: sort_key(n, last.get("司琴")),
        )
        if not piano_cands:
            raise RuntimeError(f"No valid 司琴 for Sunday {i} ({sunday})")
        piano = piano_cands[0]

        # 投影
        proj_cands = sorted(
            get_proj_pool(i, mc, proj_pool, hard_pairs, no_same_day,
                          unavailable, sundays),
            key=lambda n: sort_key(n, last.get("投影")),
        )
        if not proj_cands:
            raise RuntimeError(f"No valid 投影 for Sunday {i} ({sunday})")
        proj = proj_cands[0]

        scores[mc]    += 1
        scores[piano] += 1
        scores[proj]  += 1
        role_counts[(mc, "司會")]    += 1
        role_counts[(piano, "司琴")] += 1
        role_counts[(proj, "投影")]  += 1

        schedule.append({"date": sunday, "司會": mc, "司琴": piano, "投影": proj})

    return schedule


def local_swap(schedule, all_people, cfg_tables, sunday_months, sundays,
               max_passes=5) -> list[dict]:
    improved = True
    passes = 0
    while improved and passes < max_passes:
        improved = False
        passes += 1
        for role in ("司會", "司琴", "投影"):
            for i in range(len(schedule)):
                for j in range(i + 1, len(schedule)):
                    if schedule[i][role] == schedule[j][role]:
                        continue
                    candidate = deepcopy(schedule)
                    candidate[i][role], candidate[j][role] = (
                        candidate[j][role], candidate[i][role]
                    )
                    if not is_valid_full(candidate, cfg_tables, sunday_months, sundays):
                        continue
                    old_var = score_variance(compute_scores(schedule, all_people))
                    new_var = score_variance(compute_scores(candidate, all_people))
                    if new_var < old_var - 1e-9:
                        schedule = candidate
                        improved = True
    return schedule


def solve_schedule(sundays, mc_pool, piano_pool, proj_pool,
                   all_people, cfg_tables, sunday_months,
                   num_seeds=30) -> list[dict]:
    best_schedule = None
    best_var = float("inf")
    for seed in range(num_seeds):
        try:
            sched = greedy_pass(seed, sundays, mc_pool, piano_pool, proj_pool,
                                all_people, cfg_tables, sunday_months)
            sched = local_swap(sched, all_people, cfg_tables, sunday_months, sundays)
            var = score_variance(compute_scores(sched, all_people))
            if var < best_var:
                best_var = var
                best_schedule = sched
        except RuntimeError:
            continue
    if best_schedule is None:
        raise RuntimeError("Could not find any valid schedule. Check constraints.")
    return best_schedule


# ─────────────────────────────────────────────────────────────────────────────
# Terminal output
# ─────────────────────────────────────────────────────────────────────────────

def print_schedule(schedule, scores, cfg, sundays, mc_pool, piano_pool, proj_pool,
                   max_per_role, avail_months) -> None:
    year = cfg["year"]
    quarter = cfg["quarter"]
    q_names = {1: "第一季", 2: "第二季", 3: "第三季", 4: "第四季"}
    sep = "=" * 66

    start_d = f"{sundays[0].month}/{sundays[0].day}"
    end_d   = f"{sundays[-1].month}/{sundays[-1].day}"

    print(sep)
    print(f"  客語堂 Q{quarter} {year} 主日服事表（{year}/{start_d} – {year}/{end_d}）")
    print(sep)

    pair_mc = set(cfg["constraints"].get("hard_pairs_lookup", {}).keys())
    hard_pairs_lookup = {
        item["mc"]: item["projection_must_be"]
        for item in cfg["constraints"].get("hard_pairs", [])
    }
    pair_mc = set(hard_pairs_lookup.keys())

    def fmt(name):
        return f"{name}({scores[name]})"

    col_w = 14
    print(f"  {'日期':<10}{'司會(分)':^{col_w}}{'司琴(分)':^{col_w}}{'投影(分)'}")
    print("  " + "─" * 8 + "  " + "─" * 12 + "  " + "─" * 12 + "  " + "─" * 12)

    for day in schedule:
        date_str = f"{day['date'].month}/{day['date'].day}"
        mc, piano, proj = day["司會"], day["司琴"], day["投影"]
        note = " ★" if mc in pair_mc else ""
        print(f"  {date_str:<10}{fmt(mc):^{col_w}}{fmt(piano):^{col_w}}{fmt(proj)}{note}")

    print()
    print("  ★ = 優先配對已執行 (MC-投影 hard pair enforced)")
    print()

    mc_counts    = defaultdict(int)
    piano_counts = defaultdict(int)
    proj_counts  = defaultdict(int)
    for day in schedule:
        mc_counts[day["司會"]]    += 1
        piano_counts[day["司琴"]] += 1
        proj_counts[day["投影"]]  += 1

    print(sep)
    print("  服事次數統計")
    print(sep)

    print("\n  【司會】")
    for p in mc_pool:
        cross = ""
        if piano_counts.get(p, 0) > 0:
            cross = f"  （另司琴 {piano_counts[p]} 次）"
        elif proj_counts.get(p, 0) > 0:
            cross = f"  （另投影 {proj_counts[p]} 次）"
        print(f"    {p}: {mc_counts.get(p, 0)}{cross}")

    print("\n  【司琴】")
    for p in piano_pool:
        extras = []
        cap = max_per_role.get((p, "司琴"))
        if cap is not None:
            extras.append(f"上限 {cap} 次")
        key = (p, "司琴")
        if key in avail_months and avail_months[key]:
            months_str = "、".join(f"{m}月" for m in sorted(avail_months[key]))
            extras.append(f"僅{months_str}")
        if mc_counts.get(p, 0) > 0:
            extras.append(f"另司會 {mc_counts[p]} 次")
        note = f"  （{'，'.join(extras)}）" if extras else ""
        print(f"    {p}: {piano_counts.get(p, 0)}{note}")

    print("\n  【投影】")
    for p in proj_pool:
        cross = ""
        if mc_counts.get(p, 0) > 0:
            cross = f"  （另司會 {mc_counts[p]} 次）"
        print(f"    {p}: {proj_counts.get(p, 0)}{cross}")

    print()
    print(sep)
    print("  服事總分（含跨角色）")
    print(sep)

    vals = list(scores.values())
    col = 0
    line = "  "
    for p in list(dict.fromkeys(mc_pool + piano_pool + proj_pool)):
        token = f"{p}: {scores[p]}  "
        if col + len(token) > 62:
            print(line)
            line = "  "
            col = 0
        line += token
        col += len(token)
    if line.strip():
        print(line)

    print()
    print(f"  分差 (max − min): {max(vals)} − {min(vals)} = {max(vals) - min(vals)}")
    print(sep)


# ─────────────────────────────────────────────────────────────────────────────
# Excel export
# ─────────────────────────────────────────────────────────────────────────────

def export_xlsx(schedule, scores, cfg, sundays, output_path) -> None:
    year    = cfg["year"]
    quarter = cfg["quarter"]
    q_names = {1: "第一季", 2: "第二季", 3: "第三季", 4: "第四季"}

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "服事表1"

    N        = len(sundays)
    last_col = 2 + N

    ws.column_dimensions["A"].width = 3.69
    ws.column_dimensions["B"].width = 12.46
    for ci in range(3, last_col + 1):
        ws.column_dimensions[get_column_letter(ci)].width = 10.0

    ws.row_dimensions[1].height = 24.55
    ws.row_dimensions[2].height = 25.30
    ws.row_dimensions[3].height = 21.00
    ws.row_dimensions[4].height = 30.00
    ws.row_dimensions[5].height = 30.00
    ws.row_dimensions[6].height = 30.00

    _center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    _border = Border(
        left=Side(border_style="thin"), right=Side(border_style="thin"),
        top=Side(border_style="thin"),  bottom=Side(border_style="thin"),
    )
    _fill = PatternFill(patternType="solid", fgColor=Color(theme=0, type="theme"))

    def style(cell, bold=True, size=10.0, fill=False, border=False):
        cell.font      = Font(bold=bold, size=size)
        cell.alignment = _center
        if fill:   cell.fill   = _fill
        if border: cell.border = _border

    # Row 1: title
    tc = ws.cell(row=1, column=6)
    tc.value     = f"              {year} 年 {q_names[quarter]} 主日服事表"
    tc.font      = Font(bold=True, size=12)
    tc.alignment = Alignment(vertical="center")

    # Row 2: dates
    ws.cell(row=2, column=2).value = " 類別   日期"
    style(ws.cell(row=2, column=2), fill=True, border=True)
    for i, sunday in enumerate(sundays):
        cell = ws.cell(row=2, column=3 + i)
        cell.value         = datetime.datetime(sunday.year, sunday.month, sunday.day)
        cell.number_format = "m/d;@"
        style(cell, fill=True, border=True)

    # Row 3: 講員
    ws.cell(row=3, column=2).value = "講員"
    style(ws.cell(row=3, column=2), fill=True, border=True)
    for i in range(N):
        cell = ws.cell(row=3, column=3 + i)
        cell.value = ""
        style(cell, fill=True, border=True)

    # Rows 4-6: roles
    for row_idx, section_char, role_label, role_key in [
        (4, "客", "司會",        "司會"),
        (5, "語", "司琴",        "司琴"),
        (6, "堂", "投影製作放映", "投影"),
    ]:
        ws.cell(row=row_idx, column=1).value = section_char
        style(ws.cell(row=row_idx, column=1))
        ws.cell(row=row_idx, column=2).value = role_label
        style(ws.cell(row=row_idx, column=2))
        for i, day in enumerate(schedule):
            cell = ws.cell(row=row_idx, column=3 + i)
            name = day[role_key]
            cell.value = f"{name}\n({scores[name]})"
            style(cell)

    wb.save(output_path)
    print(f"\n  已儲存 Excel 檔案: {os.path.basename(output_path)}")


# ─────────────────────────────────────────────────────────────────────────────
# Main
# ─────────────────────────────────────────────────────────────────────────────

def main():
    config_path = sys.argv[1] if len(sys.argv) > 1 else os.path.join(BASE_DIR, "config.json")
    cfg = load_config(config_path)

    year    = cfg["year"]
    quarter = cfg["quarter"]
    mc_pool    = cfg["roles"]["司會"]
    piano_pool = cfg["roles"]["司琴"]
    proj_pool  = cfg["roles"]["投影"]
    all_people = list(dict.fromkeys(mc_pool + piano_pool + proj_pool))

    sundays = quarter_sundays(year, quarter, cfg.get("include_first_sunday", True))
    cfg_tables = build_tables(cfg, sundays)
    max_per_role, avail_months, hard_pairs, no_same_day, sunday_months, unavailable = cfg_tables

    print("正在計算最佳排班...")
    schedule = solve_schedule(sundays, mc_pool, piano_pool, proj_pool,
                              all_people, cfg_tables, sunday_months)
    scores = compute_scores(schedule, all_people)

    print_schedule(schedule, scores, cfg, sundays, mc_pool, piano_pool, proj_pool,
                   max_per_role, avail_months)

    q_names = {1: "第一季", 2: "第二季", 3: "第三季", 4: "第四季"}
    output_path = os.path.join(BASE_DIR, f"恩霖堂{year}{q_names[quarter]}主日服事表_客語堂.xlsx")
    export_xlsx(schedule, scores, cfg, sundays, output_path)


if __name__ == "__main__":
    main()
